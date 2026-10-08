"""tabular.py: 실무 서식 표(제목 행·병합 머리글·가로 펼침·합계 행·문서 안의 표) → 제출값 변환."""
import io, pandas as pd, pytest, tabular, compare
from pathlib import Path
from openpyxl import Workbook
S = Path(__file__).resolve().parents[1] / "sample_data"

class F:
    def __init__(self, name, data): self.name, self._d = name, data
    def getvalue(self): return self._d

def xlsx(rows, merges=(), sheet="Sheet1"):
    wb = Workbook(); ws = wb.active; ws.title = sheet
    for r in rows: ws.append(r)
    for m in merges: ws.merge_cells(m)
    b = io.BytesIO(); wb.save(b); return F("실적.xlsx", b.getvalue())

WIDE = [["자기주도학습센터 개소 실적 (2026. 8. 30. 기준)"], [],
        ["시군구", "개소", None, "등록 학생 수", "등원율(%)", "비고"], [None, "운영", "준비", None, None, None],
        ["가군", 2, 0, 150, 67.8, None], ["나군", 1, 1, 90, "71.2%", "사후 보정"], ["다군", 1, 0, "1,234", 59.0, None],
        ["합계", 4, 1, 1474, None, None], [], ["※ 등원율은 월 등원일수 ÷ 운영일수 × 100"]]

def test_wide_sheet_title_merged_header_totals():
    g = tabular.read_grid(xlsx(WIDE, ["A1:F1", "A3:A4", "B3:C3"], "개소 실적(시군구별)")); info = tabular.analyze(g)
    assert info.shape == "wide" and info.title.startswith("자기주도학습센터") and info.base_date_hint == "2026-08-30"
    assert info.header == ["시군구", "개소 운영", "개소 준비", "등록 학생 수", "등원율(%)", "비고"]
    assert info.center_col == 0 and info.value_cols == [1, 2, 3, 4] and [g.row_offset + r for r in info.total_rows] == [8]
    assert info.notes == ["※ 등원율은 월 등원일수 ÷ 운영일수 × 100"]
    df = tabular.wide_to_long(g, info, 0, [3, 4], info.base_date_hint)
    assert df["indicator"].unique().tolist() == ["등록 학생 수", "등원율"]          # 사전에 있으면 정규 지표명, (%) 제거
    assert len(df) == 6 and set(df["center"]) == {"가군", "나군", "다군"}           # 합계 행 제외(3센터 × 2지표)
    r = df[(df.center == "나군") & (df.indicator == "등원율")].iloc[0]
    assert r["value"] == 71.2 and r["source_row"] == 6 and r["source_sheet"] == "개소 실적(시군구별)" and r["base_date"] == "2026-08-30"
    assert df[(df.center == "다군") & (df.indicator == "등록 학생 수")]["value"].iloc[0] == 1234.0
    with_tot = tabular.wide_to_long(g, info, 0, [3], "2026-08-30", drop_totals=False)
    assert "합계" in set(with_tot["center"])

def test_wide_custom_indicator_names_and_errors():
    g = tabular.read_grid(xlsx(WIDE)); info = tabular.analyze(g)
    df = tabular.wide_to_long(g, info, 0, [1, 2], "2026-08-30", {1: "운영 센터 수", 2: "준비 센터 수"})
    assert set(df["indicator"]) == {"운영 센터 수", "준비 센터 수"}
    with pytest.raises(ValueError, match="기준일"): tabular.wide_to_long(g, info, 0, [1], "")
    with pytest.raises(ValueError, match="값 열"): tabular.wide_to_long(g, info, 0, [], "2026-08-30")
    with pytest.raises(ValueError, match="필수 컬럼 없음.*가로 펼침"): compare.load_values(xlsx(WIDE))

def test_long_sheet_with_title_row_still_loads():
    rows = [["등원율 제출본 (2026. 6. 30. 기준)"], [], ["지표명", "센터명", "기준일", "값"], ["등원율", "센터A", "2026-06-30", 60.5], ["등원율", "센터B", "2026. 6. 30.", "61%"]]
    df = compare.load_values(xlsx(rows))
    assert df["source_row"].tolist() == [4, 5] and df["value"].tolist() == [60.5, 61.0] and df["base_date"].tolist() == ["2026-06-30"] * 2

def test_year_header_row_detected():
    rows = [["시군구", 2024, 2025, 2026], ["가군", 1, 2, 3], ["나군", 2, 2, 4]]
    info = tabular.analyze(tabular.read_grid(xlsx(rows)))
    assert info.header == ["시군구", "2024", "2025", "2026"] and info.shape == "wide" and len(info.data_rows) == 2

def test_sample_wide_file_matches_long_sample():
    g = tabular.read_grid(S / "실적표_가로형_2026-06-30기준.xlsx"); info = tabular.analyze(g)
    df = tabular.wide_to_long(g, info, info.center_col, info.value_cols, info.base_date_hint)
    m = compare.compare(df[df.indicator == "등원율"], compare.load_values(S / "등원율_2026-06-30기준_7월제출본.xlsx"))
    assert (m["판정"] == "일치").sum() == 12

def test_tables_from_document_text_pipe_and_whitespace():
    doc = ("수신: ○○○ 의원실\n제목: 자기주도학습센터 운영 현황 회신\n1. 2026. 6. 30. 기준 센터별 등원율은 아래와 같습니다.\n"
           "센터명 | 등록 학생 수 | 등원율(%)\n센터A | 120 | 67.8\n센터B | 95 | 71.2\n합계 | 215 | -\n"
           "2. 끝.\n\n[PDF식 공백 표]\n구분    2024    2025\n가군    1    2\n나군    2    3\n")
    grids = tabular.grids_from_text(doc, "회신.hwpx")
    assert [g.source_sheet for g in grids] == ["표 1", "표 2"] and grids[0].row_offset == 4 and grids[1].kind == "document"
    info = tabular.analyze(grids[0])
    assert info.shape == "wide" and info.center_col == 0 and info.value_cols == [1, 2] and info.base_date_hint is None   # 기준일은 본문에 있어 담당자가 입력
    df = tabular.wide_to_long(grids[0], info, 0, [1, 2], "2026-06-30", source_version="회신 문서(회신.hwpx) 표에서 추출")
    assert len(df) == 4 and df["source_version"].iloc[0].startswith("회신 문서") and df["source_file"].iloc[0] == "회신.hwpx" and df["source_row"].tolist() == [5, 5, 6, 6]
    info2 = tabular.analyze(grids[1]); assert info2.header == ["구분", "2024", "2025"] and len(info2.data_rows) == 2
    assert tabular.grids_from_text("표 없음\n한 줄\n") == []

def test_from_records_manual_entry():
    man = pd.DataFrame([{"지표명": "등원율", "센터명": "센터A", "기준일": "2026-06-30", "값": "67.8"}, {"지표명": None, "센터명": None, "기준일": None, "값": None}], columns=tabular.MANUAL_COLS)
    df = tabular.from_records(man)
    assert len(df) == 1 and df["value"][0] == 67.8 and df["source_file"][0] == "직접 입력"
    with pytest.raises(ValueError, match="빈 행"): tabular.from_records(pd.DataFrame([{"지표명": "등원율", "센터명": None, "기준일": "2026-06-30", "값": 1}], columns=tabular.MANUAL_COLS))

def test_suggest_mapping_uses_claude_and_validates(monkeypatch):
    import llm, providers
    class Fake(providers.AnthropicProvider):
        def __init__(s): super().__init__({"api_key": "sk-test"}); s.calls = []
        def create_message(s, **kw):
            s.calls.append(kw)
            class U: input_tokens = 500; output_tokens = 80
            class B: type = "text"; text = '{"center_col": 0, "value_cols": [{"col": 4, "indicator": "등원율"}, {"col": 9, "indicator": "없는 열"}], "base_date": "2026. 8. 30.", "note": "비고 열은 값이 아님"}'
            class M: content = [B()]; usage = U(); stop_reason = "end_turn"
            return M()
    fake = Fake(); llm.reset(); monkeypatch.setattr(llm, "_provider", lambda: fake)
    g = tabular.read_grid(xlsx(WIDE)); info = tabular.analyze(g)
    out = tabular.suggest_mapping(g, info)
    assert out == {"center_col": 0, "value_cols": [(4, "등원율")], "base_date": "2026-08-30", "note": "비고 열은 값이 아님"}
    sent = fake.calls[0]["messages"][0]["content"]
    assert "등원율(%)" in sent and "가군" in sent and fake.calls[0]["schema"] is not None
    llm.reset()

FX = Path(__file__).parent / "fixtures"

def test_hwpx_merged_header_table_grid_and_conversion():
    """한글식 표: 세로·가로 병합 머리글(cellAddr·cellSpan) → 격자 복원 → 머리글 합침 → 변환."""
    import docread
    data = (FX / "회신_병합머리글.hwpx").read_bytes()
    tables = docread.hwpx_tables(data)
    assert len(tables) == 1 and tables[0][0] == ["구분", "개소", "개소", "등록 학생 수", "등원율(%)"] and tables[0][1][1:3] == ["운영", "준비"]
    grids, text = tabular.grids_from_document("회신_병합머리글.hwpx", data)
    assert len(grids) == 1 and "2026. 6. 30. 기준" in text and grids[0].kind == "document"
    info = tabular.analyze(grids[0])
    assert info.header == ["구분", "개소 운영", "개소 준비", "등록 학생 수", "등원율(%)"] and info.center_col == 0 and info.value_cols == [1, 2, 3, 4]
    df = tabular.wide_to_long(grids[0], info, 0, [3, 4], "2026-06-30", source_version="회신 문서 표에서 추출")
    assert len(df) == 6 and set(df["indicator"]) == {"등록 학생 수", "등원율"} and df["source_row"].tolist() == [3, 3, 4, 4, 5, 5]

def test_hwpx_nested_frame_table_picks_inner_data_table():
    """공문 틀(바깥 표) 안에 내용 표가 중첩된 경우: 두 표 모두 나오고, 숫자가 많은 안쪽 표가 기본 선택. 바깥 셀 글자에 안쪽 표 글자가 섞이지 않음."""
    import docread
    data = (FX / "회신_공문틀_중첩.hwpx").read_bytes()
    tables = docread.hwpx_tables(data)
    assert len(tables) == 2 and tables[0][0] == ["○○공사 ○○센터"] and tables[0][1] == [None] and tables[1][2][0] == "센터A"
    grids, text = tabular.grids_from_document("회신_공문틀_중첩.hwpx", data)
    assert [g.source_sheet for g in grids] == ["표 1"] and grids[0].rows[2][0] == "센터A"       # 1열짜리 틀 표는 제외
    assert text.count("센터A") == 1 and "○○공사 ○○센터" in text
    assert tabular.numeric_cells(grids[0]) >= 12

def test_docx_merged_header_table():
    data = (FX / "회신_병합머리글.docx").read_bytes()
    grids, _ = tabular.grids_from_document("회신_병합머리글.docx", data)
    info = tabular.analyze(grids[0])
    assert info.header == ["센터명", "지표 등록 학생 수", "지표 등원율(%)"] and info.center_col == 0 and info.value_cols == [1, 2]
    df = tabular.wide_to_long(grids[0], info, 0, [1, 2], "2026-06-30", {1: "등록 학생 수", 2: "등원율"})
    assert df["value"].tolist() == [120.0, 67.8, 95.0, 71.2]

def test_numbers_with_korean_units():
    assert tabular.is_num("120명") and tabular.is_num("3개소") and tabular.is_num("1,234천원") and not tabular.is_num("센터A")
    assert compare.clean_number("120명") == 120.0 and compare.clean_number("1,234천원") == 1234.0 and compare.clean_number("67.8 %") == 67.8
    g = tabular.read_grid(xlsx([["센터명", "등록 학생 수"], ["센터A", "120명"], ["센터B", "95명"]])); info = tabular.analyze(g)
    assert info.value_cols == [1] and tabular.wide_to_long(g, info, 0, [1], "2026-06-30")["value"].tolist() == [120.0, 95.0]

def test_hwp5_legacy_file_is_converted_and_tables_found():
    """HWP 5.0(구형식) 회신 공문: 메모리 안에서 HWPX로 변환해 본문·표를 읽는다(공문 틀 중첩 + 병합 머리글 픽스처의 HWP 판)."""
    import docread
    data = (FX / "회신_공문틀_중첩.hwp").read_bytes()
    assert data[:8] == b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1"                           # OLE 복합문서 서명
    text = docread.read("회신.hwp", data)
    assert "2026. 6. 30. 기준" in text and "센터A | 1 | 0 | 120 | 67.8" in text
    grids, _ = tabular.grids_from_document("회신.hwp", data)
    assert len(grids) == 1 and grids[0].source_file == "회신.hwpx"
    info = tabular.analyze(grids[0])
    assert info.header == ["구분", "개소 운영", "개소 준비", "등록 학생 수", "등원율(%)"] and info.center_col == 0
    df = tabular.wide_to_long(grids[0], info, 0, [4], "2026-06-30")
    assert df["value"].tolist() == [67.8, 71.2, 66.4]
    with pytest.raises(ValueError, match="변환하지 못했습니다"): docread.read("x.hwp", b"not an hwp file")

def test_center_column_prefers_most_distinct_text_column():
    """시·도(세로 병합으로 반복) | 시·군·구 | 값… 표에서는 고유값이 많은 시·군·구가 센터 열."""
    rows = [["시·도", "시·군·구", "개소 수"], ["경기", "가군", 2], ["경기", "나군", 1], ["강원", "다군", 1], ["합계", "", 4]]
    info = tabular.analyze(tabular.read_grid(xlsx(rows)))
    assert info.center_col == 1 and info.value_cols == [2]

def test_duplicate_row_names_combine_label_columns():
    """여러 시·도에 같은 구 이름이 있으면 시·도 + 시·군·구를 이어 붙여 센터명으로. 연도 × 항목 표도 연도 + 항목."""
    rows = [["시·도", "시·군·구", "개소 수"], ["서울", "중구", 2], ["부산", "중구", 1], ["부산", "남구", 1], ["합계", "", 4]]
    g = tabular.read_grid(xlsx(rows)); info = tabular.analyze(g)
    assert info.center_cols == [0, 1] and info.value_cols == [2]
    df = tabular.wide_to_long(g, info, info.center_cols, [2], "2026-06-30")
    assert df["center"].tolist() == ["서울 중구", "부산 중구", "부산 남구"]
    rows = [["연도", "세부항목", "편성액", "집행액"], [2025, "인건비", 100, 90], [2025, "운영비", 50, 40], [2026, "인건비", 120, 30], [2026, "운영비", 60, 10]]
    g = tabular.read_grid(xlsx(rows)); info = tabular.analyze(g)
    assert info.center_cols == [0, 1] and info.value_cols == [2, 3]
    df = tabular.wide_to_long(g, info, info.center_cols, info.value_cols, "2026-06-30")
    assert len(df) == 8 and "2025 인건비" in set(df["center"])
    with pytest.raises(ValueError, match="센터"): tabular.wide_to_long(g, info, [], [2], "2026-06-30")

def test_same_named_value_columns_get_column_suffix():
    rows = [["구분", "편성액", "집행액", "편성액", "집행액"], ["센터A", 1, 2, 3, 4]]
    g = tabular.read_grid(xlsx(rows)); info = tabular.analyze(g)
    df = tabular.wide_to_long(g, info, [0], [1, 2, 3, 4], "2026-06-30")
    assert df["indicator"].tolist() == ["편성액 (2열)", "집행액 (3열)", "편성액 (4열)", "집행액 (5열)"]

def test_sparse_note_column_is_not_center():
    rows = [["연도", "세부항목", "편성액", "비고"], [2025, "인건비", 100, None], [2025, "운영비", 50, "이월"], [2026, "인건비", 120, None], [2026, "운영비", 60, None]]
    info = tabular.analyze(tabular.read_grid(xlsx(rows)))
    assert info.center_col == 1 and info.center_cols == [0, 1] and info.value_cols == [2]


# ---- .xls 세 가지: 구형 바이너리(xlrd), 이름만 .xls인 HTML 표, 탭 구분 텍스트 ----
class _Up:
    def __init__(s, name, data): s.name, s._d = name, data
    def getvalue(s): return s._d

LONG_ROWS = [["지표명", "센터명", "기준일", "값"], ["등원율", "센터A", "2026-06-30", 60.3], ["등원율", "센터B", "2026-06-30", 59.4]]

def test_xls_binary_via_xlrd():
    xlwt = pytest.importorskip("xlwt")
    import io as _io
    wb = xlwt.Workbook(encoding="utf-8"); ws = wb.add_sheet("월별")
    for i, r in enumerate(LONG_ROWS):
        for j, v in enumerate(r): ws.write(i, j, v)
    buf = _io.BytesIO(); wb.save(buf); data = buf.getvalue()
    up = _Up("센터별월간출결_261008.xls", data)
    assert compare.sheet_kind(up.name, data) == "xls" and compare.list_sheets(up) == ["월별"]
    df = compare.load_values(up); assert len(df) == 2 and df["value"].tolist() == [60.3, 59.4] and df["source_sheet"].iloc[0] == "월별"

def test_xls_actually_html_table_cp949():
    html = "<html><body><table><tr><td>지표명</td><td>센터명</td><td>기준일</td><td>값</td></tr>" + "".join(
        f"<tr><td>{r[0]}</td><td>{r[1]}</td><td>{r[2]}</td><td>{r[3]}</td></tr>" for r in LONG_ROWS[1:]) + "</table></body></html>"
    data = html.encode("cp949"); up = _Up("월별관리인원_261008.xls", data)
    assert compare.sheet_kind(up.name, data) == "html" and compare.list_sheets(up) == []      # 표가 하나면 시트 선택 없음
    df = compare.load_values(up); assert len(df) == 2 and df["center"].tolist() == ["센터A", "센터B"]
    two = (html.replace("</table>", "</table><table><tr><td>a</td></tr></table>")).encode("cp949")
    assert compare.list_sheets(_Up("x.xls", two)) == ["표 1", "표 2"]
    g = tabular.read_grid(_Up("x.xls", two), "표 2"); assert g.rows[0][0] == "a" and g.source_sheet == "표 2"

def test_xls_actually_tab_text():
    data = "\n".join("\t".join(str(v) for v in r) for r in LONG_ROWS).encode("cp949"); up = _Up("월별상담횟수_261008.xls", data)
    assert compare.sheet_kind(up.name, data) == "text" and compare.list_sheets(up) == []
    df = compare.load_values(up); assert len(df) == 2 and df["value"].tolist() == [60.3, 59.4]
