"""요구서·회신 문서 본문 추출: txt / hwp / hwpx / pdf / docx → 문단 텍스트(표는 '셀 | 셀' 형태). HWP(5.0)는 HWPX로 변환해 읽는다.
표는 격자(grid)로도 꺼낸다(tabular가 제출값 변환에 사용).

HWPX는 정규식이 아니라 XML 파서로 읽는다. 한글이 저장한 실제 공문은
- 표가 문단(hp:p → hp:run) 안에 들어 있고, 공문 틀 자체가 큰 표이며 그 셀 안에 내용 표가 중첩되는 경우가 많다
- 병합 셀은 hp:cellAddr(행·열 위치)과 hp:cellSpan(가로·세로 병합 수)으로 표현되어 행마다 셀 개수가 다르다
그래서 셀 주소·병합 수로 격자를 복원하고(병합 셀은 그 범위에 같은 글자를 채움), 중첩 표는 바깥 표 셀 글자에서 제외한다.
접두어(hp:)에 의존하지 않도록 태그의 로컬 이름으로 판별한다."""
import io, re, zipfile
import xml.etree.ElementTree as ET

def _local(tag: str) -> str: return tag.rsplit("}", 1)[-1]

def _children(el, name): return [c for c in el if _local(c.tag) == name]

def _text_no_tbl(el) -> str:
    """요소 안 글자(hp:t)를 문단 단위로 이어 붙인다. 중첩 표(tbl) 안의 글자는 제외."""
    parts = []
    def walk(e, buf):
        for ch in e:
            ln = _local(ch.tag)
            if ln == "tbl": continue
            if ln == "t": buf.append("".join(ch.itertext()))
            elif ln == "p":
                b = []; walk(ch, b)
                t = re.sub(r"\s+", " ", "".join(b)).strip()
                if t: parts.append(t)
            else: walk(ch, buf)
    top = []; walk(el, top)
    t = re.sub(r"\s+", " ", "".join(top)).strip()
    if t: parts.insert(0, t)
    return " ".join(parts)

def _tbl_grid(tbl) -> list[list]:
    """hp:tbl → 격자. 셀 주소·병합 수를 반영하고 병합 범위에 같은 글자를 채운다."""
    cells, r_i = [], 0
    for tr in _children(tbl, "tr"):
        c_i = 0
        for tc in _children(tr, "tc"):
            addr = next(iter(_children(tc, "cellAddr")), None); span = next(iter(_children(tc, "cellSpan")), None)
            r = int(addr.get("rowAddr", r_i)) if addr is not None else r_i
            c = int(addr.get("colAddr", c_i)) if addr is not None else c_i
            cs = max(1, int(span.get("colSpan", 1))) if span is not None else 1
            rs = max(1, int(span.get("rowSpan", 1))) if span is not None else 1
            cells.append((r, c, cs, rs, _text_no_tbl(tc) or None)); c_i = c + cs
        r_i += 1
    if not cells: return []
    nrows = max(int(tbl.get("rowCnt") or 0), max(r + rs for r, _, _, rs, _ in cells))
    ncols = max(int(tbl.get("colCnt") or 0), max(c + cs for _, c, cs, _, _ in cells))
    grid = [[None] * ncols for _ in range(nrows)]
    for r, c, cs, rs, txt in cells:
        for dr in range(rs):
            for dc in range(cs):
                if r + dr < nrows and c + dc < ncols and grid[r + dr][c + dc] is None: grid[r + dr][c + dc] = txt
    return grid

def _hwpx_sections(data: bytes):
    z = zipfile.ZipFile(io.BytesIO(data))
    names = sorted(n for n in z.namelist() if n.lower().startswith("contents/section") and n.lower().endswith(".xml"))
    if not names: raise ValueError("HWPX 안에 본문(Contents/section*.xml)이 없습니다. 배포용(암호화) 문서이거나 손상된 파일일 수 있습니다.")
    for n in names:
        yield ET.fromstring(z.read(n))

def hwpx_tables(data: bytes) -> list[list[list]]:
    """문서 순서대로 모든 표(중첩 표 포함)를 격자로. 공문 틀 표 → 그 안의 내용 표 순."""
    out = []
    for root in _hwpx_sections(data):
        for el in root.iter():
            if _local(el.tag) == "tbl":
                g = _tbl_grid(el)
                if g: out.append(g)
    return out

def _hwpx_text(data: bytes) -> str:
    out = []
    for root in _hwpx_sections(data):
        in_tbl = {id(p) for el in root.iter() if _local(el.tag) == "tbl" for p in el.iter() if _local(p.tag) == "p"}
        for el in root.iter():
            ln = _local(el.tag)
            if ln == "p" and id(el) not in in_tbl:
                t = _text_no_tbl(el)
                if t: out.append(t)
            elif ln == "tbl":
                for row in _tbl_grid(el):
                    out.append(" | ".join((c or "") for c in row))
    return "\n".join(x for x in out if x.strip())

def _pdf_text(data: bytes) -> str:
    from pypdf import PdfReader
    r = PdfReader(io.BytesIO(data))
    return "\n".join((pg.extract_text() or "") for pg in r.pages)

def docx_tables(data: bytes) -> list[list[list]]:
    """docx 표 → 격자. python-docx는 병합 셀을 같은 셀 객체로 반복해 주므로 그대로 같은 글자가 채워진다."""
    from docx import Document
    d = Document(io.BytesIO(data))
    return [[[(c.text.strip() or None) for c in r.cells] for r in t.rows] for t in d.tables]

def _docx_text(data: bytes) -> str:
    from docx import Document
    d = Document(io.BytesIO(data))
    out = [p.text for p in d.paragraphs if p.text.strip()]
    for t in d.tables:
        for r in t.rows:
            out.append(" | ".join(c.text.strip() for c in r.cells))
    return "\n".join(out)

def hwp_to_hwpx(data: bytes) -> bytes:
    """HWP 5.0(구형식, OLE 복합문서) → HWPX 바이트. python-hwpx의 변환기를 쓴다(메모리 안, 외부 전송 없음).
    암호·배포용 문서는 변환기가 거부한다."""
    try:
        from hwpx.hwp5 import package
    except ImportError as e:
        raise ValueError("HWP(구형식)를 읽으려면 python-hwpx 모듈이 필요합니다(start.py가 자동 설치). 또는 한글에서 HWPX로 저장해 올리세요.") from e
    try:
        conv = package.convert(data)
        return package.to_hwpx_bytes(conv.files)
    except Exception as e:
        raise ValueError(f"HWP를 HWPX로 변환하지 못했습니다({type(e).__name__}: {str(e)[:120]}). 암호·배포용 문서이거나 HWP 3.0 이하일 수 있습니다. 한글에서 HWPX로 저장해 올리세요.") from e

def read(name: str, data: bytes) -> str:
    n = name.lower()
    if n.endswith(".hwp"): return _hwpx_text(hwp_to_hwpx(data))
    if n.endswith(".hwpx"): return _hwpx_text(data)
    if n.endswith(".pdf"): return _pdf_text(data)
    if n.endswith(".docx"): return _docx_text(data)
    return decode_text(data)

def decode_text(data: bytes) -> str:
    """txt 인코딩 판별: UTF-16 BOM → utf-8-sig → cp949(한글 Windows 기본) → euc-kr → utf-8(대체 문자)."""
    if data[:2] in (b"\xff\xfe", b"\xfe\xff"): return data.decode("utf-16", errors="ignore")
    for enc in ("utf-8-sig", "cp949", "euc-kr"):
        try: return data.decode(enc)
        except UnicodeDecodeError: continue
    return data.decode("utf-8", errors="replace")

def tables(name: str, data: bytes) -> list[list[list]] | None:
    """문서 형식별 표 격자. hwp·hwpx·docx는 표 개체에서 직접, 그 외(pdf·txt)는 None(본문 텍스트에서 찾아야 함)."""
    n = name.lower()
    if n.endswith(".hwp"): return hwpx_tables(hwp_to_hwpx(data))
    if n.endswith(".hwpx"): return hwpx_tables(data)
    if n.endswith(".docx"): return docx_tables(data)
    return None
